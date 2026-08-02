"""Targeted Reader regressions across parser, store, bridge, and playback JS."""

import base64
import io
import json
from pathlib import Path
import re
import shutil
import subprocess

import pytest

import reader_parser
import reader_store
import settings
import webui_shell


APP_JS = Path(__file__).with_name("webui") / "app.js"
APP_DIR = Path(__file__).parent
MAINTAINED_APP_DIRS = (
    APP_DIR,
    APP_DIR / "Ports" / "macOS" / "app",
    APP_DIR / "Ports" / "Linux" / "app",
)
READER_SPEECH_CONTRACT_FILES = tuple(
    platform_dir / relative
    for platform_dir in MAINTAINED_APP_DIRS
    for relative in (
        Path("webui/app.js"),
        Path("webui_shell.py"),
        Path("ai/tts_providers.py"),
        Path("ai/__init__.py"),
        Path("test_tts_providers.py"),
        Path("test_reader.py"),
    )
)


@pytest.fixture
def isolated_reader(monkeypatch, tmp_path):
    monkeypatch.setattr(reader_store, "PATH", str(tmp_path / "reader_library.json"))
    return tmp_path


def test_structured_speech_text_matches_persisted_positions(isolated_reader):
    parsed = reader_parser.ParserRegistry.parse_bytes(
        b"# Title\n\n2. **Second** item\n3. [Third](https://example.com)",
        "sample.md",
    )
    assert parsed.plain_text() == "Title\n\nSecond item\n\nThird"

    doc_id = reader_store.save_doc_parsed(parsed, fmt="markdown")
    opened = reader_store.get_doc(doc_id)
    assert opened["text"] == parsed.plain_text()
    assert reader_store.list_docs()[0]["length"] == 4

    assert reader_store.set_position(doc_id, 999) is True
    assert reader_store.list_docs()[0]["percent"] == 100
    assert reader_store.continue_reading_doc() is None


def test_legacy_structured_row_self_heals_without_duplicate(isolated_reader):
    blocks = [
        {"type": "heading", "text": "Title", "level": 1},
        {"type": "list_item", "text": "Item", "level": 1},
    ]
    assert reader_store._save([{
        "id": "legacy", "title": "Legacy", "text": "# Title\n\n• Item",
        "length": 4, "position": 4, "opened": 1, "added": 1,
        "blocks": blocks, "format": "markdown", "reading_sessions": [],
    }])

    assert reader_store.get_doc("legacy")["text"] == "Title\n\nItem"
    assert reader_store.list_docs()[0]["length"] == 2
    parsed = reader_parser.ParsedDocument(
        title="Legacy",
        blocks=[reader_parser.ContentBlock.from_dict(block) for block in blocks],
    )
    new_id = reader_store.save_doc_parsed(parsed, fmt="markdown")
    assert new_id != "legacy"
    assert len(reader_store._load()) == 1


def test_library_save_never_evicts_old_documents(isolated_reader):
    rows = [{
        "id": str(index), "title": str(index), "text": "one word",
        "length": 2, "position": 0, "opened": index, "added": index,
    } for index in range(125)]
    assert reader_store._save(rows) is True
    assert len(reader_store._load()) == 125


def test_session_detail_cap_preserves_cumulative_totals(isolated_reader):
    doc_id = reader_store.save_doc("Doc", "one two three")
    for _ in range(27):
        assert reader_store.log_reading_session(doc_id, 0, 2, 60) is True

    raw = reader_store._load()[0]
    summary = reader_store.list_docs()[0]
    assert len(raw["reading_sessions"]) == 25
    assert summary["session_count"] == 27
    assert summary["total_minutes"] == 27.0


def test_empty_collections_are_real_and_survive_last_document_removal(isolated_reader):
    assert reader_store.create_collection("Research") is True
    assert reader_store.list_collections() == [{"name": "Research", "count": 0}]
    doc_id = reader_store.save_doc("Doc", "one two")
    assert reader_store.add_to_collection(doc_id, "Research") is True
    assert reader_store.list_collections()[0]["count"] == 1
    assert reader_store.remove_from_collection(doc_id, "Research") is True
    assert reader_store.list_collections() == [{"name": "Research", "count": 0}]
    assert reader_store.delete_collection("Research") == 0
    assert reader_store.list_collections() == []


def test_corrupt_library_recovers_backup_and_preserves_damaged_bytes(isolated_reader):
    doc_id = reader_store.save_doc("Doc", "one two three")
    assert reader_store.set_position(doc_id, 1) is True  # creates .bak
    Path(reader_store.PATH).write_text("{broken", encoding="utf-8")

    recovered = reader_store._load()
    assert recovered and recovered[0]["id"] == doc_id
    assert json.loads(Path(reader_store.PATH).read_text(encoding="utf-8"))[0]["id"] == doc_id
    assert list(Path(reader_store.PATH).parent.glob("reader_library.json.corrupt-*"))

    # A crash after rotating main -> .bak but before tmp -> main leaves main
    # absent; startup must restore the backup rather than show an empty library.
    assert reader_store.set_position(doc_id, 2) is True
    Path(reader_store.PATH).unlink()
    assert reader_store._load()[0]["id"] == doc_id


def test_text_html_markdown_and_rtf_import_fidelity():
    utf16 = reader_parser.ParserRegistry.parse_bytes(
        "Hello smart reader".encode("utf-16"), "utf16.txt")
    assert utf16.plain_text() == "Hello smart reader"
    assert reader_parser.ParserRegistry.parse_bytes(
        b"extensionless text", "README").plain_text() == "extensionless text"
    assert "." not in reader_parser.ParserRegistry.supported_extensions()

    html = reader_parser.ParserRegistry.parse_bytes(
        b'<meta charset="windows-1252"><blockquote><p>caf\xe9</p></blockquote>',
        "legacy.html",
    )
    assert html.plain_text() == "café"
    assert len(html.blocks) == 1

    md = reader_parser.ParserRegistry.parse_bytes(
        b"| **Name** | Value |\n| --- | ---: |\n| [Alpha](https://x) | 1 |",
        "table.md",
    )
    assert md.blocks[0].type == "table"
    assert md.blocks[0].meta["headers"] == ["Name", "Value"]
    assert md.blocks[0].meta["rows"] == [["Alpha", "1"]]

    semicolon = reader_parser.ParserRegistry.parse_bytes(
        b"Name;Value\nAlpha;1\n", "semicolon.csv")
    assert semicolon.blocks[0].meta["headers"] == ["Name", "Value"]

    rtf = reader_parser._rtf_to_text(
        rb"{\rtf1{\fonttbl{\f0 Arial;}} hello {\*\unknown SECRET}"
        rb" A\u8217\'92B can\rquote t \emdash \bullet item}")
    assert "Arial" not in rtf and "SECRET" not in rtf
    assert "A’B" in rtf and "can’t" in rtf and "—" in rtf and "• item" in rtf


def test_docx_tables_remain_in_document_order():
    from docx import Document

    document = Document()
    document.add_paragraph("Before")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Alpha"
    table.cell(1, 1).text = "42"
    document.add_paragraph("After")
    buf = io.BytesIO()
    document.save(buf)

    parsed = reader_parser.ParserRegistry.parse_bytes(buf.getvalue(), "table.docx")
    assert [block.type for block in parsed.blocks] == ["paragraph", "table", "paragraph"]
    assert parsed.blocks[1].meta["rows"] == [["Alpha", "42"]]


def test_odt_preserves_spacing_and_nested_list_levels():
    from odf.opendocument import OpenDocumentText
    from odf.text import LineBreak, List, ListItem, P, S, Tab

    document = OpenDocumentText()
    para = P()
    para.addText("Alpha")
    para.addElement(S(c=2))
    para.addText("Beta")
    para.addElement(Tab())
    para.addText("Gamma")
    para.addElement(LineBreak())
    para.addText("Delta")
    document.text.addElement(para)
    outer = List()
    parent = ListItem()
    parent.addElement(P(text="Parent"))
    inner = List()
    child = ListItem()
    child.addElement(P(text="Child"))
    inner.addElement(child)
    parent.addElement(inner)
    outer.addElement(parent)
    document.text.addElement(outer)
    buf = io.BytesIO()
    document.save(buf)

    parsed = reader_parser.ParserRegistry.parse_bytes(buf.getvalue(), "nested.odt")
    assert parsed.blocks[0].text == "Alpha  Beta\tGamma\nDelta"
    items = [block for block in parsed.blocks if block.type == "list_item"]
    assert [(item.text, item.level) for item in items] == [("Parent", 1), ("Child", 2)]


def test_bridge_enforces_file_and_extracted_text_limits(monkeypatch, tmp_path):
    api = webui_shell.Api.__new__(webui_shell.Api)
    monkeypatch.setattr(webui_shell, "MAX_READER_IMPORT_BYTES", 4)
    oversized = tmp_path / "large.txt"
    oversized.write_bytes(b"12345")
    assert "larger than" in api.reader_import_file(str(oversized))["message"]

    empty = api.reader_import_bytes(base64.b64encode(b"").decode(), "empty.txt")
    assert empty["ok"] is False
    assert "No readable text" in empty["message"]

    monkeypatch.setattr(webui_shell, "MAX_READER_TEXT_CHARS", 4)
    too_much = api.reader_save("", "12345")
    assert too_much["ok"] is False


def test_bridge_propagates_stats_failure_and_does_not_persist_fallback():
    api = webui_shell.Api.__new__(webui_shell.Api)

    class FailedStats:
        def record_reader_session(self, **_kwargs):
            return False

    api._stats = lambda: FailedStats()
    assert api.reader_record_stats(10, 5)["ok"] is False

    class MemorySettings:
        def __init__(self):
            self.values = {
                "pro_mode": True,
                "local_only_mode": False,
                "instant_text": False,
                "reader_tts_provider": "openrouter",
                "reader_tts_model": "google/gemini-3.1-flash-tts-preview",
                "openrouter_api_key": "sk-or-test",
            }
            self.updates = []

        def get(self, key, default=None):
            return self.values.get(key, default)

        def update(self, **values):
            self.updates.append(values)

    api.settings = MemorySettings()
    original = webui_shell.ai.synthesize_with_fallback
    webui_shell.ai.synthesize_with_fallback = lambda *_args, **_kwargs: (
        b"audio", "audio/mpeg", {
            "ok": True, "provider": "openrouter",
            "model": "mistralai/voxtral-mini-tts-2603",
            "voice": "gb_oliver_neutral", "fallback": True,
            "fallback_provider": "openrouter",
        })
    try:
        result = api.reader_tts("hello", provider="openrouter")
    finally:
        webui_shell.ai.synthesize_with_fallback = original
    assert result["ok"] is True and result["fallback"] is True
    assert api.settings.updates == []


def test_reader_fetch_deduplicates_and_stale_request_cannot_fill_new_cache():
    node = shutil.which("node")
    if not node:
        pytest.skip("Node.js is unavailable")
    source = APP_JS.read_text(encoding="utf-8")
    match = re.search(
        r"^const MAX_READER_AUDIO_CACHE_CHUNKS.*?(?=^function readerTick)",
        source, flags=re.MULTILINE | re.DOTALL)
    assert match
    script = match.group(0) + r"""
const READER = {
  chunks: [{text: "hello"}], cache: {}, fetching: {}, chunkIdx: 0,
  model: "model", provider: "provider", effectiveModel: null,
  effectiveVoice: null, effectiveProvider: null, synthTimes: [],
  synthDone: 0, synthTotal: 1, fallbackUsed: false
};
let calls = 0;
let call = async () => {
  calls++;
  await new Promise(resolve => setTimeout(resolve, 15));
  return {ok: true, mime: "audio/mpeg", audio: "YQ=="};
};
function readerCurrentVoice() { return "voice"; }
function readerUpdateSynthProgress() {}
function toast() {}
(async () => {
  await Promise.all([readerFetchChunk(0), readerFetchChunk(0)]);
  if (calls !== 1) throw new Error("duplicate synthesis: " + calls);

  READER.cache = {};
  READER.fetching = {};
  let release;
  call = () => new Promise(resolve => { release = resolve; });
  const old = readerFetchChunk(0);
  const replacement = {};
  READER.cache = replacement;
  READER.fetching = {};
  release({ok: true, mime: "audio/mpeg", audio: "Yg=="});
  await old;
  if (Object.keys(replacement).length) throw new Error("stale cache write");
})().catch(error => { console.error(error); process.exit(1); });
"""
    result = subprocess.run(
        [node, "-e", script], capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr


def test_reader_ui_has_real_cancellation_failure_and_collection_paths():
    source = APP_JS.read_text(encoding="utf-8")
    open_doc = source[source.index("async function readerOpenDoc"):
                      source.index("async function readerAddAndRead")]
    assert open_doc.index("READER.audio.pause()") < open_doc.index(
        'await call("reader_open", id)')
    assert "if (openGen !== READER.gen) return;" in open_doc
    assert "readerBuildBlocks(doc.blocks, doc.text)" in open_doc

    play_chunk = source[source.index("async function readerPlayChunk"):
                        source.index("async function readerPlay()")]
    assert "Never skip unread text" in play_chunk
    catch_end = play_chunk.index("// Prefetch the next two chunks")
    assert "return;" in play_chunk[play_chunk.index("} catch (e)"):catch_end]
    assert "readerStopListeningClock();" in play_chunk

    assert 'call("reader_create_collection", name)' in source
    assert 'await call("reader_list_collections")' in source


def test_reader_sync_is_opt_in_and_disclosure_names_exact_speech_attempt():
    assert settings.DEFAULTS["sync_reader"] is False
    source = APP_JS.read_text(encoding="utf-8")
    disclosure = source[source.index("async function confirmReaderCloudUse"):
                        source.index("function readerBuildPane")]
    assert (
        "Mumble freezes that provider and the selected model for this request, "
        "makes one synthesis attempt with that exact pair, and fails closed "
        "without a local, sibling-model, or cross-provider fallback"
    ) in disclosure
    assert (
        "another compatible model from that same provider "
        "may be tried"
    ) not in disclosure
    assert "Reader Sync, when enabled" in disclosure
    cloud = Path(__file__).with_name("cloud_sync.py").read_text(encoding="utf-8")
    assert 'payload.pop("source_path", None)' in cloud


def test_reader_speech_docs_match_frozen_attempt_on_maintained_platforms():
    stale_fragments = (
        "another compatible model from that same provider " + "may be tried",
        "On failure, may try a compatible model " + "from the same provider",
        "Same-provider model fallback via " + "synthesize_with_fallback()",
        "and same-provider model " + "fallback.",
        "It may fall back to a " + "sibling model",
        "a sibling model " + "must use",
    )
    offenders = {}
    for path in READER_SPEECH_CONTRACT_FILES:
        source = path.read_text(encoding="utf-8")
        found = [fragment for fragment in stale_fragments if fragment in source]
        if found:
            offenders[str(path.relative_to(APP_DIR))] = found
    assert offenders == {}

    for platform_dir in MAINTAINED_APP_DIRS:
        bridge = (platform_dir / "webui_shell.py").read_text(encoding="utf-8")
        bridge_words = " ".join(bridge.split())
        assert (
            "The frozen decision authorizes exactly one synthesis attempt "
            "with the selected provider and model."
        ) in bridge_words

        provider_docs = (platform_dir / "ai" / "tts_providers.py").read_text(
            encoding="utf-8")
        assert (
            "One exact frozen Reader provider/model attempt via "
            "synthesize_with_fallback()"
        ) in provider_docs

        provider_tests = (platform_dir / "test_tts_providers.py").read_text(
            encoding="utf-8")
        assert "one exact frozen Reader provider/model attempt." in provider_tests
